from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from smoothment.statement.balance import check_balances
from smoothment.statement.model import StatementFormatError
from smoothment.statement.parsers import bbva
from tests.conftest import FIXTURES


def test_english_row_count_and_first_row() -> None:
    statement = bbva.parse_current(FIXTURES / "bbva_en.xlsx", "Main", None)
    assert statement.bank == "bbva"
    assert statement.kind == "current"
    assert len(statement.rows) == 14
    first = statement.rows[0]
    assert first.date == datetime(2025, 12, 5)
    assert first.amount == Decimal("-95.00")
    assert first.balance == Decimal("108.79")
    assert first.counterparty == "John Smith store examplecityname    es"
    assert first.line == 6
    for row in statement.rows:
        assert isinstance(row.amount, Decimal)


def test_english_transfer_row() -> None:
    statement = bbva.parse_current(FIXTURES / "bbva_en.xlsx", "Main", None)
    row = statement.rows[1]
    assert row.counterparty is None
    assert row.bank_type == "Sent from revolut"
    assert row.reference == "Sent from Revolut"


def test_english_chain_and_balances_check_out() -> None:
    statement = bbva.parse_current(FIXTURES / "bbva_en.xlsx", "Main", None)
    assert len(statement.chains) == 1
    chain = statement.chains[0]
    assert chain.name == "EUR"
    assert chain.indexes == tuple(range(13, -1, -1))
    assert check_balances(statement) == ()


def _write_spanish_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Informe BBVA"
    sheet["D2"] = "Ultimos movimientos"
    for _ in range(2):
        sheet.append([])
    sheet.append(
        [
            None,
            "F.Valor",
            "Fecha",
            "Concepto",
            "Movimiento",
            "Importe",
            "Divisa",
            "Disponible",
            "Divisa",
            "Observaciones",
        ]
    )
    sheet.append(
        [
            None,
            "04/09/2026",
            "04/09/2026",
            "Transferencia recibida",
            "Sent from revolut",
            550,
            "EUR",
            569.19,
            "EUR",
            "Sent from Revolut",
        ]
    )
    sheet.append(
        [
            None,
            "31/08/2026",
            "31/08/2026",
            "Some payee",
            "Some movement",
            -545.58,
            "EUR",
            19.19,
            "EUR",
            "Some comment",
        ]
    )
    sheet.append(
        [
            None,
            "25/07/2026",
            "28/07/2026",
            "Sample Taxi",
            "Pago con tarjeta",
            -10.9,
            "EUR",
            564.77,
            "EUR",
            "SAMPLE TAXI",
        ]
    )
    workbook.save(path)


def test_spanish_workbook_parses_day_first_dates(tmp_path: Path) -> None:
    path = tmp_path / "bbva_es.xlsx"
    _write_spanish_workbook(path)

    statement = bbva.parse_current(path, "Main", None)

    assert len(statement.rows) == 3
    assert statement.rows[0].date == datetime(2026, 9, 4)
    assert statement.rows[2].date == datetime(2026, 7, 25)
    assert statement.rows[2].counterparty == "Sample Taxi"
    assert statement.rows[0].counterparty is None
    assert check_balances(statement) == ()


def test_unknown_header_raises(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for _ in range(4):
        sheet.append([])
    sheet.append([None, "Wrong", "Header", "Row", "Here", "X", "Y", "Z", "W", "V"])
    path = tmp_path / "bbva_bad_header.xlsx"
    workbook.save(path)

    with pytest.raises(StatementFormatError) as exc:
        bbva.parse_current(path, "Main", None)
    assert exc.value.row == 5
    assert exc.value.column == "header"


def test_bad_date_cell_raises(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for _ in range(4):
        sheet.append([])
    sheet.append(
        [
            None,
            "Eff. Date",
            "Date",
            "Item",
            "Transaction",
            "Amount",
            "Foreign currency",
            "Available",
            "Foreign currency",
            "Comments",
        ]
    )
    sheet.append(
        [None, "soon", "12/05/2025", "Payee", "Card payment", -10, "EUR", 100, "EUR", "note"]
    )
    path = tmp_path / "bbva_bad_date.xlsx"
    workbook.save(path)

    with pytest.raises(StatementFormatError) as exc:
        bbva.parse_current(path, "Main", None)
    assert exc.value.row == 6


def test_input_rows_counts_non_empty_sheet_rows_after_header() -> None:
    statement = bbva.parse_current(FIXTURES / "bbva_en.xlsx", "Main", None)
    assert statement.input_rows == 14
