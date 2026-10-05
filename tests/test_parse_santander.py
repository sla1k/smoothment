from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from smoothment.statement.balance import check_balances
from smoothment.statement.model import StatementFormatError
from smoothment.statement.parsers import santander
from tests.conftest import FIXTURES


def test_iban_reads_fixture() -> None:
    assert santander.iban(FIXTURES / "santander.xlsx") == "ES0000000000000000000000"


def test_row_count_and_first_row() -> None:
    statement = santander.parse_current(FIXTURES / "santander.xlsx", "Main", None)
    assert statement.bank == "santander"
    assert statement.kind == "current"
    assert len(statement.rows) == 4
    first = statement.rows[0]
    assert first.line == 9
    assert first.date == datetime(2025, 12, 5)
    assert first.amount == Decimal("-43.15")
    assert first.balance == Decimal("56.85")
    for row in statement.rows:
        assert isinstance(row.amount, Decimal)


def test_chain_and_balances_check_out() -> None:
    statement = santander.parse_current(FIXTURES / "santander.xlsx", "Main", None)
    assert len(statement.chains) == 1
    chain = statement.chains[0]
    assert chain.name == "EUR"
    assert chain.indexes == (3, 2, 1, 0)
    assert check_balances(statement) == ()


def test_evidence_from_concepto() -> None:
    statement = santander.parse_current(FIXTURES / "santander.xlsx", "Main", None)
    rows = statement.rows

    assert rows[0].counterparty == "Acme Restaurant"
    assert rows[0].bank_type == "Pago Movil"
    assert rows[0].reference is None

    assert rows[1].counterparty == "John Doe"
    assert rows[1].bank_type == "Transferencia"
    assert rows[1].reference == "Sent From Revolut"

    assert rows[2].counterparty == "John Doe"
    assert rows[2].bank_type == "Transferencia"
    assert rows[2].reference is None

    assert rows[3].counterparty == "John Doe"
    assert rows[3].bank_type == "Transferencia"
    assert rows[3].reference is None


def _write_generated_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["C1"] = santander.TITLE
    sheet.append([])
    sheet.append(["FECHA OPERACIÓN", "FECHA VALOR", "CONCEPTO", "IMPORTE EUR", "SALDO"])
    sheet.append(["05/09/2026", "05/09/2026", "Some payee", "-149.87", "500.13"])
    sheet.append(["04/09/2026", "04/09/2026", "Other payee", "0", "650"])
    workbook.save(path)


def test_generated_workbook_parses_text_amount_cells(tmp_path: Path) -> None:
    path = tmp_path / "santander_generated.xlsx"
    _write_generated_workbook(path)

    statement = santander.parse_current(path, "Main", None)

    assert len(statement.rows) == 2
    assert statement.rows[0].amount == Decimal("-149.87")
    assert statement.rows[1].amount == Decimal("0")
    assert santander.iban(path) is None


def test_missing_header_raises(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Wrong", "Header", "Row", "Here", "Too"])
    path = tmp_path / "santander_bad_header.xlsx"
    workbook.save(path)

    with pytest.raises(StatementFormatError) as exc:
        santander.parse_current(path, "Main", None)
    assert exc.value.column == "header"


def test_bad_date_cell_raises(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["FECHA OPERACIÓN", "FECHA VALOR", "CONCEPTO", "IMPORTE EUR", "SALDO"])
    sheet.append(["soon", "05/09/2026", "Some payee", "-10", "100"])
    path = tmp_path / "santander_bad_date.xlsx"
    workbook.save(path)

    with pytest.raises(StatementFormatError) as exc:
        santander.parse_current(path, "Main", None)
    assert exc.value.row == 2


def test_input_rows_counts_non_empty_sheet_rows_after_header() -> None:
    statement = santander.parse_current(FIXTURES / "santander.xlsx", "Main", None)
    assert statement.input_rows == 4


def test_same_day_rows_newest_first_have_a_consistent_chain(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["FECHA OPERACIÓN", "FECHA VALOR", "CONCEPTO", "IMPORTE EUR", "SALDO"])
    sheet.append(["05/09/2026", "05/09/2026", "Later payee", -10, 90])
    sheet.append(["05/09/2026", "05/09/2026", "Earlier payee", 100, 100])
    path = tmp_path / "santander_same_day.xlsx"
    workbook.save(path)

    statement = santander.parse_current(path, "Main", None)

    assert statement.chains[0].indexes == (1, 0)
    assert check_balances(statement) == ()
